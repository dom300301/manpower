#!/usr/bin/env python3
"""
官等員額配置比率　概念驗證（PoC）

目的：證明「用總表的編制表明細，依 114.12.4 修正後的配置準則第 5、6 條算出
　　　簡任比率、委任比率，再對照附表三」這件事做得到，並看看結果長什麼樣。

依據（grade-check/regs/）：
  第 5 條  委任比率 = 委任員額 / 各官等員額總數（下限）
           簡任比率 = 簡任員額 / 薦任員額（上限）
           只算「定有官等職等」的員額（警察、軍職、師級醫事等不計）
  第 6 條  跨列官等：
           「委任至薦任」「薦任至簡任」→ 全數以高官等計（第 1 款）
           「列委任，內 N 人得列薦任」→ N 人算薦任，其餘算委任（第 2 款）
           「委任或薦任」「薦任或簡任」→ 1/3 算低官等、2/3 算高官等，尾數算高官等（第 3、4 款）

本 PoC 的簡化（正式版要補）：
  * 機關類別只用「層級＋名稱關鍵字」粗分，附表三的工程、社福、矯正、航空站……沒有細分。
  * 「得列」備註只抓「內Ｎ人得列薦任／簡任」這種寫法。
  * 雙軌制機關、駐外機構、編制員額未滿 12 人的機關沒有排除。

用法：  python3 grade-check/poc/ratio_poc.py
輸出：  grade-check/poc/官等比率_PoC.csv
"""
import collections
import csv
import os
import re
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
MASTER = os.path.join(ROOT, 'org-compare', 'data', '組織法編制表_總表_20261031.xlsx')
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '官等比率_PoC.csv')

CN = {'一': 1, '二': 2, '三': 3, '四': 4, '五': 5, '六': 6, '七': 7, '八': 8, '九': 9}


def c2n(s):
    """中文數字（一～九十九）→ 整數。"""
    s = s.strip()
    if s.isdigit():
        return int(s)
    if '十' in s:
        a, b = s.split('十', 1)
        return (CN.get(a, 1) if a else 1) * 10 + (CN.get(b, 0) if b else 0)
    return CN.get(s, 0)


def split_grades(grade, n, remark):
    """把一列編制表的員額拆成 委任/薦任/簡任 三個數字。回傳 None 表示不計入（非一般文職）。"""
    d = {'委任': 0, '薦任': 0, '簡任': 0}
    m = re.search(r'內([一二三四五六七八九十\d]+)人得列(簡任|薦任)', remark)
    up = c2n(m.group(1)) if m else 0
    up_grade = m.group(2) if m else None
    if grade in d:                       # 單列一個官等
        d[grade] += n - up
        if up_grade:
            d[up_grade] += up            # 第 6 條第 2 款
    elif grade in ('委任或薦任', '委任至薦任', '薦任或簡任', '薦任至簡任'):
        lo, hi = re.split('或|至', grade)
        if '至' in grade:                # 第 6 條第 1 款：全以高官等計
            d[hi] += n
        else:                            # 第 6 條第 3、4 款：1/3 低、2/3 高，尾數歸高
            lo_n = n // 3
            d[lo] += lo_n
            d[hi] += n - lo_n
    else:
        return None
    return d


def category(name, level):
    """粗分機關類別 →（類別名, 簡任上限%, 委任下限%）。正式版要改成人工確認的對照表。"""
    research = re.search('研究|試驗|實驗|改良|繁殖', name)
    if level == '2':
        return '中央二級機關', 40, 10
    if level == '3':
        if research:
            return '中央研究機關', 35, 15
        if re.search('署$|局$', name):
            return '中央三級_署局', 35, 10
        if re.search('訓練|學院', name):
            return '中央三級_訓練機關', 35, 10
        return '中央三級_行政機關', 35, 15
    if level == '4':
        if research:
            return '中央研究機關', 35, 15
        return '中央四級_行政機關(粗分)', 20, 20
    return '未分類', None, None


def main():
    wb = openpyxl.load_workbook(MASTER, read_only=True)
    ov = list(wb['機關總覽'].iter_rows(values_only=True))
    oi = {h: i for i, h in enumerate(ov[0])}
    level_of = {r[oi['機關']]: str(r[oi['層級']]) for r in ov[1:]}
    det = list(wb['編制表明細'].iter_rows(values_only=True))
    di = {h: i for i, h in enumerate(det[0])}

    by_agency = collections.defaultdict(list)
    for r in det[1:]:
        by_agency[r[di['機關名稱']]].append(r)

    rows = []
    for name, lines in by_agency.items():
        W = J = K = total = skipped = 0
        for r in lines:
            try:
                n = int(r[di['員額']])
            except (TypeError, ValueError):
                continue
            if n <= 0:
                continue
            d = split_grades((r[di['官等']] or '').strip(), n, r[di['備註']] or '')
            if d is None:
                skipped += n
                continue
            W += d['委任']; J += d['薦任']; K += d['簡任']; total += n
        if total == 0:
            continue
        cat, j_lim, w_lim = category(name, level_of.get(name, '?'))
        sj = round(K / J * 100, 1) if J else None
        sw = round(W / total * 100, 1)
        flags = []
        if j_lim is not None and sj is not None and sj > j_lim:
            flags.append(f'簡任比率 {sj}% 超過上限 {j_lim}%')
        if w_lim is not None and sw < w_lim:
            flags.append(f'委任比率 {sw}% 低於下限 {w_lim}%')
        if total < 12:
            flags.append('員額未滿12人，不適用附表三')
        rows.append([name, level_of.get(name, ''), cat, total, skipped, W, J, K, sj, j_lim, sw, w_lim, '；'.join(flags)])

    with open(OUT, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(['機關', '層級', '類別(粗分)', '定有官等員額', '未計入員額(警/軍/師級等)',
                    '委任', '薦任', '簡任', '簡任/薦任%', '簡任上限%', '委任/總數%', '委任下限%', '初判'])
        w.writerows(sorted(rows, key=lambda x: (x[1], x[2], x[0])))

    cnt = collections.Counter()
    for r in rows:
        cnt[r[2]] += 1
        if '簡任比率' in r[12]:
            cnt[r[2] + '|簡任超上限'] += 1
        if '委任比率' in r[12]:
            cnt[r[2] + '|委任低於下限'] += 1
    print(f'已寫出 {OUT}（{len(rows)} 個機關）')
    for c in sorted(k for k in cnt if '|' not in k):
        print(f'{c}: {cnt[c]} 機關，簡任超上限 {cnt[c + "|簡任超上限"]}，委任低於下限 {cnt[c + "|委任低於下限"]}')


if __name__ == '__main__':
    sys.exit(main())
