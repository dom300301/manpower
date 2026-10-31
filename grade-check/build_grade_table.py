#!/usr/bin/env python3
"""
以總表（585 個行政院所屬機關的編制表明細）建立「職稱列等對照表（實證版）」。

想法：職務列等表的現行各表不在公開網站上，但每個機關的編制表都經考試院核備，
把同層級、同職稱在各機關實際訂列的官等職等統計起來，就是「現行列等」最完整、
可追溯（每筆都有法規資料庫 PDF）的依據。

產出：
  grade-check/data/職稱列等對照表_實證_<日期>.xlsx   人看的（總覽、列等組合明細、機關類型細分、職稱清單、說明）
  grade-check/rules/title_grades.json                程式用（規則引擎、網頁）

用法：
  python3 grade-check/build_grade_table.py [--master org-compare/data/組織法編制表_總表_20261031.xlsx]
"""
import argparse
import collections
import datetime as dt
import json
import os
import re
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

CN = {'一': 1, '二': 2, '三': 3, '四': 4, '五': 5, '六': 6, '七': 7, '八': 8, '九': 9}


def c2n(s):
    """中文數字（一～九十九）或阿拉伯數字 → 整數。"""
    s = (s or '').strip()
    if s.isdigit():
        return int(s)
    if '十' in s:
        a, b = s.split('十', 1)
        return (CN.get(a, 1) if a else 1) * 10 + (CN.get(b, 0) if b else 0)
    return CN.get(s, 0)


def parse_grades(s):
    """'第八職等至第九職等' → [8, 9]；'第五職等或第六職等至第七職等' → [5, 6, 7]。回傳排序後的職等清單。"""
    nums = [c2n(x) for x in re.findall(r'第([一二三四五六七八九十]+)職等', s or '')]
    if not nums:
        return []
    out = set()
    # 「至」連成區間，「或」分段
    parts = re.split(r'或', s)
    for p in parts:
        ns = [c2n(x) for x in re.findall(r'第([一二三四五六七八九十]+)職等', p)]
        if not ns:
            continue
        if '至' in p and len(ns) >= 2:
            out.update(range(min(ns), max(ns) + 1))
        else:
            out.update(ns)
    return sorted(out)


def grade_code(guan, zhi):
    """把官等＋職等濃縮成短碼，例如 '薦任 8-9'、'委任或薦任 5|6-7'、'簡任 10-11'。"""
    guan = (guan or '').strip()
    zhi = (zhi or '').strip()
    if not guan and not zhi:
        return ''
    seg = []
    for p in re.split(r'或', zhi):
        ns = [c2n(x) for x in re.findall(r'第([一二三四五六七八九十]+)職等', p)]
        if not ns:
            continue
        seg.append(f'{min(ns)}-{max(ns)}' if ('至' in p and len(ns) >= 2 and min(ns) != max(ns)) else str(ns[0]))
    z = '|'.join(seg) if seg else zhi
    return f'{guan} {z}'.strip()


DE_LIE = re.compile(r'內([一二三四五六七八九十\d]+)人得列(簡任|薦任)第([一二三四五六七八九十]+)職等(?:至第([一二三四五六七八九十]+)職等)?')


def parse_remark(b):
    b = (b or '').strip()
    f = {
        '暫列': '暫列' in b,
        '法律所定': '組織法律所定' in b,
        '兼任': bool(re.search(r'由.{1,30}兼任', b)),
        '聘任': ('聘任' in b) or ('教育人員任用條例' in b),
        '醫事': bool(re.search(r'師（[一二三]）級|師\([一二三]\)級|醫事人員', b)),
        '得列': [],
    }
    for m in DE_LIE.finditer(b):
        f['得列'].append({'人數': c2n(m.group(1)), '官等': m.group(2), '職等': f'{c2n(m.group(3))}' + (f'-{c2n(m.group(4))}' if m.group(4) else '')})
    return f


AGENCY_TYPE_RE = re.compile(r'(分署|分局|分院|總隊|總臺|總處|署|局|處|所|館|中心|醫院|學院|大學|學校|中學|委員會|部|院|監獄|看守所|站|隊|園區|之家|場|廠|電臺|銀行)$')


def agency_type(name):
    m = AGENCY_TYPE_RE.search(name)
    if m:
        return m.group(1)
    if re.search(r'[北中南高基]關$', name):
        return '海關'
    return '其他'


# 附表二　訂有官等職等職稱屬性區分表（擇要；用於標示職稱類別）
TITLE_CLASS = {
    '工程類': '總工程司、副總工程司、工程司、副工程司、正工程司、主任工程司、幫工程司、工程員、助理工程員、工務員',
    '研究類': '研究委員、研究員、副研究員、助理研究員、研究助理、研析員、研譯員',
    '資訊類': '高級分析師、分析師、高級設計師、設計師、設計員、助理設計師、助理程式設計師、高級管理師、管理師、資訊管理師、助理管理師、操作員、電子作業管理師',
    '技術類': '技監、技正、技士、技佐、總技師、技師、副技師、技術員、主任管制員、管制員、助理管制員、主任報務員、報務員、主任航務員、航務員、飛行員、主任航詢員、航詢員、主任氣象員、預報員、通訊員、觀測員、助理觀測員、督察員、測量員、電務員、報務長、勞工安全管理員、勞工安全衛生管理員、安全衛生管理員、專利高級審查官、專利審查官、專利助理審查官、獸醫師、獸醫、醫用物理師、放射化學師、放射物理師、醫院放射安全師、語言訓練員、病房助理員、專任教官、艦艇駕駛員、艦艇機師、技術助理員、技術生、檢驗師、檢驗員、警報員、線務員、職業安全管理師、職業安全衛生管理員',
    '行政技術通用': '顧問、稽察、稽察員、衛生稽查員、訓練師、助理訓練師、輔導員、實習指導員、檢查員、作業員、調查專員、調查官、調查員、解說員、心理輔導員、心理測驗員、參議',
    '司法法務類': '書記官長、書記官、通譯、法警長、副法警長、法警、執達員、錄事、庭務員、主任公設辯護人、公設辯護人、主任觀護人、觀護人、主任公證人、公證人、主任檢察事務官、檢察事務官、主任技術審查官、技術審查官、主任司法事務官、司法事務官、主任行政執行官、行政執行官、執行員、主任法醫師、法醫師、典獄長、副典獄長、監長、教誨師、作業導師、主任管理員、管理員、教導員、導師、主任調查保護官、少年調查官、少年保護官、家事調查官、廉政專員、廉政官、主任商業調查官、商業調查官',
    '駐外人員類': '大使、公使、常任代表、副常任代表、副參事、經濟參事、經濟副參事、教育參事、教育副參事、文化參事、文化副參事、僑務參事、僑務副參事、科技參事、科技副參事、總領事、副總領事、領事、副領事、主事',
}
TITLE_CLASS_MAP = {}
for cls, names in TITLE_CLASS.items():
    for n in names.split('、'):
        TITLE_CLASS_MAP.setdefault(n.strip(), cls)

HEAD_RE = re.compile(r'(?<!副)(長|主任)$')


def title_class(title):
    t = title.strip()
    base = re.sub(r'[（(].*?[)）]', '', t)
    if base in TITLE_CLASS_MAP:
        return TITLE_CLASS_MAP[base]
    if re.search(r'(師|士|生)$', base) and re.search(r'醫|護|藥|治療|檢驗|放射|營養|心理|助產|牙體|聽力|呼吸|語言', base):
        return '醫事類'
    if re.search(r'(長|主任|秘書長|總工程司)$', base) or base in ('主任秘書',):
        return '主管類（副主管含）' if base.startswith('副') else '主管類'
    return '行政類'


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--master', default=os.path.join(ROOT, 'org-compare', 'data', '組織法編制表_總表_20261031.xlsx'))
    ap.add_argument('--outdir', default=os.path.join(HERE, 'data'))
    a = ap.parse_args()
    os.makedirs(a.outdir, exist_ok=True)
    os.makedirs(os.path.join(HERE, 'rules'), exist_ok=True)

    wb = openpyxl.load_workbook(a.master, read_only=True)
    ov = list(wb['機關總覽'].iter_rows(values_only=True))
    oi = {h: i for i, h in enumerate(ov[0])}
    agencies = {}
    for r in ov[1:]:
        agencies[r[oi['機關']]] = {'層級': str(r[oi['層級']]), '隸屬': r[oi['隸屬']], '部會': r[oi['部會']],
                                  '類型': agency_type(r[oi['機關']]), '編制表': r[oi['編制表來源']], '驗算': r[oi['驗算狀態']]}
    det = list(wb['編制表明細'].iter_rows(values_only=True))
    di = {h: i for i, h in enumerate(det[0])}

    rows = []
    for r in det[1:]:
        name = r[di['機關名稱']]
        ag = agencies.get(name)
        if not ag:
            continue
        title = (r[di['職稱']] or '').strip()
        if not title:
            continue
        try:
            n = int(r[di['員額']])
        except (TypeError, ValueError):
            n = 0
        guan = (r[di['官等']] or '').strip()
        zhi = (r[di['職等']] or '').strip()
        rem = parse_remark(r[di['備註']])
        rows.append({'機關': name, '層級': ag['層級'], '類型': ag['類型'], '部會': ag['部會'], '職稱': title, '單位': r[di['單位']],
                     '官等': guan, '職等': zhi, '碼': grade_code(guan, zhi), '職等數': parse_grades(zhi), '員額': n, **rem,
                     '備註': (r[di['備註']] or '').strip(), '來源': ag['編制表']})

    # ---- 彙整：層級 × 職稱 ----
    agg = {}
    for x in rows:
        k = (x['層級'], x['職稱'])
        g = agg.setdefault(k, {'機關': set(), '員額': 0, '列數': 0, '組合': {}, '暫列機關': set(), '法律所定機關': set(),
                               '兼任列': 0, '聘任列': 0, '零員額列': 0, '得列': collections.Counter(), '類型': collections.Counter(),
                               '職等': []})
        g['機關'].add(x['機關']); g['員額'] += x['員額']; g['列數'] += 1
        if x['兼任']:
            g['兼任列'] += 1
        if x['聘任']:
            g['聘任列'] += 1
        if x['員額'] <= 0:          # 員額 0 的列（多為「由○○兼任」）不計入列等分布
            g['零員額列'] += 1
            continue
        key_code = x['碼'] or '（未定官等職等）'
        c = g['組合'].setdefault(key_code, {'官等': x['官等'], '職等': x['職等'], '機關': set(), '員額': 0, '暫列機關': set(), '法律所定機關': set(), '類型': collections.Counter()})
        c['機關'].add(x['機關']); c['員額'] += x['員額']; c['類型'][x['類型']] += 1
        if x['暫列']:
            g['暫列機關'].add(x['機關']); c['暫列機關'].add(x['機關'])
        if x['法律所定']:
            g['法律所定機關'].add(x['機關']); c['法律所定機關'].add(x['機關'])
        for d in x['得列']:
            g['得列'][f"得列{d['官等']}{d['職等']}"] += 1
        g['類型'][x['類型']] += 1
        g['職等'] += x['職等數']

    # ---- Excel ----
    out = openpyxl.Workbook()
    ws0 = out.active; ws0.title = '說明'
    today = dt.date.today().isoformat()
    notes = [
        '職稱列等對照表（實證版）',
        f'產生日期：{today}　來源：{os.path.basename(a.master)}（行政院所屬機關編制表明細，資料取自全國法規資料庫）',
        '',
        '一、這張表是什麼：把同一層級、同一職稱在各機關編制表「實際訂列」的官等職等統計出來。每一筆都可追溯到該機關的編制表 PDF（機關總覽的「編制表來源」）。',
        '二、為什麼這樣做：職務列等表現行各表不在公開網站上（詳 regs/職務列等表/README.md）；編制表經考試院核備，是現行列等最完整的實證依據。',
        '三、怎麼讀：',
        '　「列等碼」＝官等＋職等，例如「薦任 8-9」＝薦任第八職等至第九職等；「委任或薦任 5|6-7」＝委任第五職等或薦任第六職等至第七職等。',
        '　「主要列等」＝採用機關數最多的列等；「全部列等組合」列出每種列等的機關數與員額。',
        '　「暫列機關數」＝備註載「本職稱之官等職等暫列」者，表示該列等尚未納入職務列等表（經考試院同意先暫列）。',
        '　「法律所定機關數」＝備註載「為組織法律所定」者，列等由組織法律直接規定，不受職務列等表拘束。',
        '　「得列」＝備註「內 N 人得列薦（簡）任第 X 職等」的統計。',
        '四、限制：',
        '　1. 機關類型由機關名稱字尾推定（署、局、處、所、館、中心、醫院、監獄…），僅供分組參考。',
        '　2. 總表「檢核_待人工確認」內的機關，其明細可能有誤；比對時請對照 PDF。',
        '　3. 114 年考試院三次通案調整（中央四級機關一級業務單位主管等）生效後尚未修編的機關，其編制表仍是舊列等；本表反映的是「已核備的現況」，不是「可列的上限」。',
        '　4. 雙軌制（警察、醫事、聘任）職稱的官等職等欄多為空白或師級，另以「未定官等職等」列出。',
        '五、工作表：職稱列等總覽（層級×職稱）、列等組合明細（層級×職稱×列等）、機關類型細分、職稱清單（含附表二屬性）。',
    ]
    for i, t in enumerate(notes, 1):
        ws0.cell(row=i, column=1, value=t)
    ws0.column_dimensions['A'].width = 120

    def style_header(ws):
        for c in ws[1]:
            c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDEBF7'); c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.freeze_panes = 'A2'

    ws1 = out.create_sheet('職稱列等總覽')
    hdr = ['層級', '職稱', '職稱屬性(附表二)', '機關數', '員額合計', '列等組合數', '主要列等', '主要列等機關數', '主要列等占比%',
           '全部列等組合(機關數/員額)', '最低職等', '最高職等', '暫列機關數', '法律所定機關數', '得列註記', '兼任列數', '聘任列數', '零員額列數', '機關類型分布', '範例機關(主要列等)']
    ws1.append(hdr)
    json_out = {}
    for (lv, title), g in sorted(agg.items(), key=lambda kv: (kv[0][0], -len(kv[1]['機關']), kv[0][1])):
        combos = sorted(g['組合'].items(), key=lambda kv: (-len(kv[1]['機關']), -kv[1]['員額']))
        if not combos:          # 全部都是兼任／零員額列
            combos = [('（僅兼任或零員額）', {'官等': '', '職等': '', '機關': set(), '員額': 0, '暫列機關': set(), '法律所定機關': set(), '類型': collections.Counter()})]
        main_code, main = combos[0]
        n_ag = len(g['機關'])
        all_txt = '；'.join(f"{code}：{len(c['機關'])}機關/{c['員額']}人" + ('（暫列' + str(len(c['暫列機關'])) + '）' if c['暫列機關'] else '') for code, c in combos)
        zs = g['職等']
        ws1.append([int(lv), title, title_class(title), n_ag, g['員額'], len(combos), main_code, len(main['機關']), round(len(main['機關']) / n_ag * 100, 1),
                    all_txt, min(zs) if zs else None, max(zs) if zs else None, len(g['暫列機關']), len(g['法律所定機關']),
                    '；'.join(f'{k}×{v}' for k, v in g['得列'].most_common()), g['兼任列'], g['聘任列'], g['零員額列'],
                    '；'.join(f'{k}{v}' for k, v in g['類型'].most_common(6)), '、'.join(sorted(main['機關'])[:3])])
        json_out.setdefault(lv, {})[title] = {
            '屬性': title_class(title), '機關數': n_ag, '員額': g['員額'], '主要列等': main_code,
            '最低職等': min(zs) if zs else None, '最高職等': max(zs) if zs else None,
            '暫列機關數': len(g['暫列機關']), '法律所定機關數': len(g['法律所定機關']),
            '得列': dict(g['得列']),
            '列等組合': [{'碼': code, '官等': c['官等'], '職等': c['職等'], '職等數': parse_grades(c['職等']), '機關數': len(c['機關']), '員額': c['員額'],
                      '暫列機關數': len(c['暫列機關']), '法律所定機關數': len(c['法律所定機關']), '機關類型': dict(c['類型'].most_common(5)),
                      '機關': sorted(c['機關'])} for code, c in combos],
        }
    style_header(ws1)
    for col, w in zip('ABCDEFGHIJKLMNOPQRST', [6, 18, 14, 8, 9, 8, 18, 9, 9, 60, 8, 8, 9, 9, 24, 8, 8, 8, 30, 40]):
        ws1.column_dimensions[col].width = w
    ws1.auto_filter.ref = ws1.dimensions

    ws2 = out.create_sheet('列等組合明細')
    ws2.append(['層級', '職稱', '列等碼', '官等', '職等', '機關數', '員額', '暫列機關數', '法律所定機關數', '機關類型', '機關清單'])
    for (lv, title), g in sorted(agg.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        for code, c in sorted(g['組合'].items(), key=lambda kv: -len(kv[1]['機關'])):
            ws2.append([int(lv), title, code, c['官等'], c['職等'], len(c['機關']), c['員額'], len(c['暫列機關']), len(c['法律所定機關']),
                        '；'.join(f'{k}{v}' for k, v in c['類型'].most_common(5)), '、'.join(sorted(c['機關']))])
    style_header(ws2)
    for col, w in zip('ABCDEFGHIJK', [6, 18, 20, 12, 30, 8, 8, 9, 9, 30, 80]):
        ws2.column_dimensions[col].width = w
    ws2.auto_filter.ref = ws2.dimensions

    # 類型細分：層級 × 類型 × 職稱 → 主要列等
    ws3 = out.create_sheet('機關類型細分')
    ws3.append(['層級', '機關類型', '職稱', '機關數', '主要列等', '主要列等機關數', '全部列等組合(機關數)'])
    agg3 = {}
    for x in rows:
        if x['員額'] <= 0:
            continue
        k = (x['層級'], x['類型'], x['職稱'])
        g = agg3.setdefault(k, {'機關': set(), '組合': collections.defaultdict(set)})
        g['機關'].add(x['機關']); g['組合'][x['碼'] or '（未定官等職等）'].add(x['機關'])
    for (lv, typ, title), g in sorted(agg3.items(), key=lambda kv: (kv[0][0], kv[0][1], -len(kv[1]['機關']), kv[0][2])):
        combos = sorted(g['組合'].items(), key=lambda kv: -len(kv[1]))
        ws3.append([int(lv), typ, title, len(g['機關']), combos[0][0], len(combos[0][1]), '；'.join(f'{c}：{len(s)}' for c, s in combos)])
    style_header(ws3)
    for col, w in zip('ABCDEFG', [6, 10, 18, 8, 18, 9, 60]):
        ws3.column_dimensions[col].width = w
    ws3.auto_filter.ref = ws3.dimensions

    ws4 = out.create_sheet('職稱清單')
    ws4.append(['職稱', '職稱屬性(附表二)', '出現層級', '機關數', '員額合計'])
    tl = {}
    for x in rows:
        t = tl.setdefault(x['職稱'], {'層級': set(), '機關': set(), '員額': 0})
        t['層級'].add(x['層級']); t['機關'].add(x['機關']); t['員額'] += x['員額']
    for title, t in sorted(tl.items(), key=lambda kv: -len(kv[1]['機關'])):
        ws4.append([title, title_class(title), '、'.join(sorted(t['層級'])), len(t['機關']), t['員額']])
    style_header(ws4)
    for col, w in zip('ABCDE', [24, 16, 10, 8, 9]):
        ws4.column_dimensions[col].width = w
    ws4.auto_filter.ref = ws4.dimensions

    xlsx = os.path.join(a.outdir, f'職稱列等對照表_實證_{today.replace("-", "")}.xlsx')
    out.save(xlsx)
    meta = {'產生日期': today, '來源': os.path.basename(a.master), '機關數': len(agencies), '明細列數': len(rows),
            '說明': '層級→職稱→列等組合（機關數、員額、暫列、法律所定、機關清單）。詳 build_grade_table.py 與 Excel 說明頁。'}
    with open(os.path.join(HERE, 'rules', 'title_grades.json'), 'w', encoding='utf-8') as f:
        json.dump({'meta': meta, 'data': json_out}, f, ensure_ascii=False, indent=1)
    print(f'已寫出 {xlsx}')
    print(f'已寫出 rules/title_grades.json（層級×職稱 {sum(len(v) for v in json_out.values())} 組）')
    return 0


if __name__ == '__main__':
    sys.exit(main())
