#!/usr/bin/env python3
"""
官職等配置審核　規則引擎（第一版）

讀總表 → 對每個機關逐條檢核「各機關職稱及官等職等員額配置準則」（114.12.4 修正）→ 輸出
  grade-check/output/check_results.json   每機關：官等計算明細、各規則結果、每列職稱列等狀態
  grade-check/output/檢核結果_全機關.csv    每機關一列的摘要
  grade-check/rules/agency_category.csv    機關類別（附表三）程式初判表（首次產生；已存在則讀取人工確認結果）

規則編號（對應 00_需求檢視與架構規劃.md 第四節）：
  R1  簡任比率 ≤ 上限（第5條、附表三；簡任÷薦任）
  R2  委任比率 ≥ 下限（第5條、附表三；委任÷各官等總數）
  R3  跨列官等拆算（第6條）——本身不是判定，而是 R1、R2 的計算基礎，會列出每列的拆算
  R3a 非主管官等職等不得高於副首長、幕僚長（第4條第1項第4款）
  R3b 各職稱官等職等應接續不得跳空（第4條第1項第6款）
  R4  委任第一至三職等職稱最低員額（第7條第1項；另有 115.9.30 刪除草案）
  R5  委四至五得列薦六員額 ≥ 委五或薦六至七職稱員額之 1/12（第7條第2項）
  R6  未滿12人機關應置委任職稱（第7條第3項）
  R7  薦任第七職等以下薦任非主管比率（第8條）
  R9  三級署局參議上限（第9條第2款）
  R10 中央二、三級簡任非主管 ≤ 二級單位主管×7/10（第9條第5款）
  R11 職稱與列等對照實證分布（rules/title_grades.json）
  R13 舊制（111.9.23 附表三）比率對照
"""
import argparse
import collections
import csv
import datetime as dt
import json
import math
import os
import re
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, 'org-compare'))

CN = {'一': 1, '二': 2, '三': 3, '四': 4, '五': 5, '六': 6, '七': 7, '八': 8, '九': 9}


def c2n(s):
    s = (s or '').strip()
    if s.isdigit():
        return int(s)
    if '十' in s:
        a, b = s.split('十', 1)
        return (CN.get(a, 1) if a else 1) * 10 + (CN.get(b, 0) if b else 0)
    return CN.get(s, 0)


def grades_of(zhi):
    """職等字串 → 涵蓋的職等集合。"""
    out = set()
    for p in re.split(r'或', zhi or ''):
        ns = [c2n(x) for x in re.findall(r'第([一二三四五六七八九十]+)職等', p)]
        if not ns:
            continue
        if '至' in p and len(ns) >= 2:
            out.update(range(min(ns), max(ns) + 1))
        else:
            out.update(ns)
    return out


def grade_code(guan, zhi):
    seg = []
    for p in re.split(r'或', zhi or ''):
        ns = [c2n(x) for x in re.findall(r'第([一二三四五六七八九十]+)職等', p)]
        if ns:
            seg.append(f'{min(ns)}-{max(ns)}' if ('至' in p and len(ns) >= 2 and min(ns) != max(ns)) else str(ns[0]))
    z = '|'.join(seg) if seg else (zhi or '')
    return f'{(guan or "").strip()} {z}'.strip()


DE_LIE = re.compile(r'內([一二三四五六七八九十\d]+)人得列(簡任|薦任)第([一二三四五六七八九十]+)職等(?:至第([一二三四五六七八九十]+)職等)?')
DUAL_RE = re.compile(r'警|校|尉|士官|師級|士（生）級|士\(生\)級|將')
MGR_RE = re.compile(r'(?:長|主任|主任委員|主席|總裁|司令|委員長|執行長|主任秘書|秘書長|總工程司|主任工程司)$')
HEAD_TITLES = re.compile(r'^(部長|署長|局長|處長|所長|館長|院長|主任|主任委員|委員長|分署長|分局長|分院長|總隊長|臺長|站長|隊長|校長|場長|廠長|典獄長|關務長|人事長|主計長|總裁|主席|司令|中心主任|監獄長|理事主席|執行長|園長|所長兼|局長兼)$')
STAFF_HEAD = ('主任秘書', '秘書長', '副秘書長', '總工程司', '主任工程司')
FIRST_LEVEL_MGR = {'2': ('司長', '處長', '主任', '室主任', '組長', '中心主任', '執行秘書'), '3': ('組長', '處長', '主任', '室主任', '中心主任', '科長'), '4': ('科長', '課長', '主任', '室主任', '股長')}
SECOND_LEVEL_MGR = {'2': ('科長',), '3': ('科長',), '4': ('股長',)}
CANYI = ('參事', '技監', '顧問', '參議', '研究委員')  # 第9條第1、2款所稱高階參贊／調節性職務，不計入第5款


def load_rules():
    r = {}
    r['ratio'] = json.load(open(os.path.join(HERE, 'rules', 'ratio_table.json'), encoding='utf-8'))
    r['titles'] = json.load(open(os.path.join(HERE, 'rules', 'title_grades.json'), encoding='utf-8'))['data']
    return r


# ---------------------------------------------------------------- 機關類別（附表三）初判
def guess_category(name, level, parent, ministry):
    n = name
    if level == '1':
        return '中央一級機關（比照二級）', '層級1'
    if re.search(r'榮民總醫院$|醫學院附設醫院$', n):
        return '公立醫療機構_臺大成大陽明交大附設醫院及臺北臺中高雄屏東榮總', '名稱含榮民總醫院'
    if re.search(r'醫院|療養院|分院$', n):
        return '公立醫療機構_其他', '名稱含醫院／療養院／分院'
    if re.search(r'警政署$|刑事警察局$', n):
        return '警察機關_警政署刑事警察局', '名稱'
    if re.search(r'警察|警政署', n):
        return '警察機關_其他', '名稱含警察'
    if re.search(r'消防署$', n):
        return '消防機關_消防署', '名稱'
    if re.search(r'消防', n):
        return '警察機關_其他', '名稱含消防（消防機關其他類，附表三未列，暫比照警察機關其他）'
    if re.search(r'研究|試驗|實驗|改良|繁殖', n):
        return '中央研究機關', '名稱含研究／試驗／實驗／改良／繁殖（附表三附註六(三)）'
    if level == '2':
        return '中央二級機關', '層級2'
    if level == '3':
        if re.search(r'工程', n):
            return '中央三級_工程機關', '名稱含工程（仍須符合附註六(一)工程類員額比例）'
        if re.search(r'訓練|學院|文官學院', n):
            return '中央三級_訓練機關', '名稱含訓練／學院'
        if re.search(r'署$|局$', n):
            return '中央三級_署局及相當中央三級獨立機關', '名稱以署／局結尾'
        if ministry in ('教育部', '文化部') or re.search(r'博物館|圖書館|科學教育館|藝術|文化|紀念堂|美術館|資訊圖書館', n):
            return '中央三級_社教文化機關', '教育部或文化部所屬／名稱含館、藝術、文化'
        return '中央三級_行政機關', '三級其他'
    # level 4
    if re.search(r'調查處$|調查站$|調查局', n):
        return '中央四級_調查機關', '名稱含調查處／站'
    if re.search(r'監獄|看守所|戒治所|矯正學校|觀護所|技能訓練所|少年輔育院|矯正署', n):
        return '中央四級_矯正機關', '法務部矯正署所屬'
    if re.search(r'工程|工務|養護|河川分署|水利署', n):
        return '中央四級_工程機關', '名稱含工程／養護／河川（仍須符合附註六(一)）'
    if re.search(r'防疫檢疫|防檢', n):
        return '中央四級_動植物防疫檢疫機關', '名稱'
    if re.search(r'航空站|飛航服務總臺', n):
        return '中央四級_航空站及飛航服務總臺', '名稱'
    if re.search(r'訓練|學院|講習', n):
        return '中央四級_訓練機關', '名稱含訓練／學院'
    if re.search(r'榮譽國民之家|之家$|教養院|敬老院|兒童之家|少年之家|養護中心|福利', n):
        return '中央四級_社會福利機關', '名稱（附表三附註六(五)）'
    if re.search(r'行政執行', n):
        return '中央四級_行政執行機關', '名稱'
    if re.search(r'氣象站|氣象雷達站|氣象測報|天文站|海象', n):
        return '中央四級_氣象測報機關', '名稱'
    if ministry in ('教育部', '文化部') or re.search(r'博物館|圖書館|社會教育館|藝術|文化|生活美學館', n):
        return '中央四級_社教文化機關', '教育部或文化部所屬／名稱'
    return '中央四級_行政機關', '四級其他'


def load_or_make_categories(agencies):
    path = os.path.join(HERE, 'rules', 'agency_category.csv')
    cats = {}
    if os.path.exists(path):
        with open(path, encoding='utf-8-sig') as f:
            for r in csv.DictReader(f):
                cats[r['機關']] = r
        print(f'讀取機關類別對照表 {path}（{len(cats)} 機關）')
    missing = [a for a in agencies if a['機關'] not in cats]
    if missing:
        for a in missing:
            c, why = guess_category(a['機關'], a['層級'], a['隸屬'], a['部會'])
            cats[a['機關']] = {'機關': a['機關'], '層級': a['層級'], '隸屬': a['隸屬'], '部會': a['部會'], '程式初判類別': c, '初判依據': why,
                               '人工確認類別': '', '雙軌制': '是' if re.search(r'警察|警政|消防|海巡|醫院|榮民總醫院|分院|療養院|國防部', a['機關']) else '',
                               '軍文併用': '是' if re.search(r'國防部|海巡|全民防衛', a['機關']) else '', '駐外': '', '例外核備紀錄': '', '備註': ''}
        with open(path, 'w', encoding='utf-8-sig', newline='') as f:
            w = csv.DictWriter(f, fieldnames=['機關', '層級', '隸屬', '部會', '程式初判類別', '初判依據', '人工確認類別', '雙軌制', '軍文併用', '駐外', '例外核備紀錄', '備註'])
            w.writeheader(); w.writerows(cats[a['機關']] for a in agencies if a['機關'] in cats)
        print(f'已產生／補入機關類別對照表 {path}（新增 {len(missing)} 機關），請同仁於「人工確認類別」欄填寫確認結果')
    return cats


# ---------------------------------------------------------------- 官等計算（第6條）
def split_row(guan, zhi, n, remark):
    """回傳 (委任, 薦任, 簡任, 說明, 計入?)。"""
    guan = (guan or '').strip(); zhi = (zhi or '').strip(); remark = remark or ''
    if n <= 0:
        return 0, 0, 0, '員額0（兼任或空白）', False
    if not guan or DUAL_RE.search(guan) or DUAL_RE.search(zhi):
        return 0, 0, 0, '非一般文職（特任、比照、警察、軍職、師級或未定官等）不計', False
    d = {'委任': 0, '薦任': 0, '簡任': 0}
    up = 0; upg = None
    for m in DE_LIE.finditer(remark):
        up += c2n(m.group(1)); upg = m.group(2)
    up = min(up, n)
    if guan in d:
        d[guan] = n - up
        if upg:
            d[upg] += up
        note = f'{guan}{n}' + (f'（內{up}人得列{upg}）' if upg else '')
    elif guan in ('委任至薦任', '薦任至簡任'):
        hi = guan[-2:]; d[hi] = n; note = f'{guan}→全以{hi}計（第6條第1款）'
    elif guan in ('委任或薦任', '薦任或簡任'):
        lo, hi = guan.split('或'); lo_n = n // 3; d[lo] = lo_n; d[hi] = n - lo_n
        note = f'{guan}{n}→{lo}{lo_n}、{hi}{n - lo_n}（第6條第3、4款：1/3、2/3，尾數歸{hi}）'
    else:
        return 0, 0, 0, f'官等「{guan}」無法判讀，不計', False
    return d['委任'], d['薦任'], d['簡任'], note, True


def pct(a, b):
    return round(a / b * 100, 1) if b else None


def finding(rule, law, item, calc, value, threshold, result, gap='', note=''):
    return {'規則': rule, '條文': law, '項目': item, '計算': calc, '數值': value, '門檻': threshold, '結果': result, '差距': gap, '說明': note}


def check_agency(a, rows, cat, rules):
    level = a['層級']
    ratio_new = rules['ratio']['new']['中央']; ratio_old = rules['ratio']['old']['中央']
    category = (cat.get('人工確認類別') or '').strip() or cat['程式初判類別']
    confirmed = bool((cat.get('人工確認類別') or '').strip())
    dual = (cat.get('雙軌制') or '') == '是'
    F = []
    # ---- 第6條拆算
    W = J = K = 0; total = 0; excluded = 0; detail = []
    for r in rows:
        w, j, k, note, counted = split_row(r['官等'], r['職等'], r['員額'], r['備註'])
        W += w; J += j; K += k
        if counted:
            total += r['員額']
        else:
            excluded += r['員額']
        r['拆算'] = {'委任': w, '薦任': j, '簡任': k, '說明': note, '計入': counted}
    counts = {'委任': W, '薦任': J, '簡任': K, '定有官等職等員額': total, '未計入員額': excluded}
    F.append(finding('R3', '第6條', '跨列官等拆算', '委任/薦任/簡任員額', f'{W}/{J}/{K}（總數 {total}，另 {excluded} 人非一般文職不計）', '—', '計算基礎', '', '每列拆算見明細'))
    # ---- R1 / R2（新制）
    key = category if category in ratio_new else ('中央二級機關' if level in ('1', '2') else None)
    th = ratio_new.get(key) if key else None
    not_apply = None
    if total < 12:
        not_apply = '編制員額未滿12人，不適用附表三（第5條第1項但書）'
    elif (cat.get('駐外') or '') == '是':
        not_apply = '駐外機構不適用附表三'
    sj = pct(K, J); sw = pct(W, total)
    if not_apply:
        F.append(finding('R1', '第5條、附表三', '簡任比率', '簡任÷薦任', f'{sj}%' if sj is not None else '—', '—', '不適用', '', not_apply))
        F.append(finding('R2', '第5條、附表三', '委任比率', '委任÷各官等總數', f'{sw}%' if sw is not None else '—', '—', '不適用', '', not_apply))
    elif th is None:
        F.append(finding('R1', '第5條、附表三', '簡任比率', '簡任÷薦任', f'{sj}%', '—', '需確認', '', f'機關類別「{category}」在附表三無對應比率，請人工指定'))
        F.append(finding('R2', '第5條、附表三', '委任比率', '委任÷各官等總數', f'{sw}%', '—', '需確認', '', '同上'))
    else:
        lim = th.get('簡任'); low = th.get('委任')
        cat_note = f'類別：{category}' + ('' if confirmed else '（程式初判，待確認）')
        if lim is None:
            F.append(finding('R1', '第5條、附表三', '簡任比率', '簡任÷薦任', f'{sj}%' if sj is not None else '—', '不訂簡任比率', '不適用', '', cat_note))
        elif sj is None:
            F.append(finding('R1', '第5條、附表三', '簡任比率', '簡任÷薦任', '—', f'≤{lim}%', '需確認', '', '薦任員額為0，無法計算；' + cat_note))
        else:
            ok = sj <= lim
            # 差距：需減少多少簡任（或增加多少薦任）才達標
            gap = ''
            if not ok:
                need_k = K - math.floor(lim / 100 * J)
                need_j = math.ceil(K / (lim / 100)) - J
                gap = f'簡任需減 {need_k} 人，或薦任需增 {need_j} 人'
            F.append(finding('R1', '第5條、附表三', '簡任比率', f'簡任 {K} ÷ 薦任 {J}', f'{sj}%', f'≤{lim}%', '符合' if ok else '不符', gap, cat_note + ('；雙軌制機關僅以單列官等職等員額計' if dual else '')))
        ok2 = sw >= low
        gap2 = ''
        if not ok2:
            need_w = math.ceil(low / 100 * total) - W
            gap2 = f'委任需增 {need_w} 人（總數不變時）'
        F.append(finding('R2', '第5條、附表三', '委任比率', f'委任 {W} ÷ 總數 {total}', f'{sw}%', f'≥{low}%', '符合' if ok2 else '不符', gap2, cat_note))
    # ---- R13 舊制
    old_key = None
    if key:
        mp = rules['ratio']['old'].get('對照', {})
        old_key = mp.get(f'{key}@{level}') or mp.get(key) or (key if key in ratio_old else None)
    tho = ratio_old.get(old_key) if old_key else None
    if tho and not not_apply:
        sj_old = pct(K, total)
        r1 = '—' if tho.get('簡任') is None else ('符合' if sj_old <= tho['簡任'] else '不符')
        r2 = '符合' if sw >= tho['委任'] else '不符'
        F.append(finding('R13', '111.9.23 附表三（舊制）', '舊制簡任／委任比率', f'簡任 {K}÷總數 {total}；委任 {W}÷總數 {total}', f'{sj_old}% / {sw}%',
                         f'≤{tho.get("簡任")}% / ≥{tho["委任"]}%', f'{r1} / {r2}', '', f'舊類別：{old_key}。114.12.4 前核備之編制表依第12條得於下次修編時再調整'))
    # ---- R4 委任一至三職等（第7條第1項）
    low13 = sum(r['拆算']['委任'] for r in rows if r['拆算']['計入'] and grades_of(r['職等']) and max(grades_of(r['職等'])) <= 3 and (r['官等'] or '').strip() == '委任')
    if total >= 30:
        need = 1 if total < 100 else 2
        F.append(finding('R4', '第7條第1項', '委任第一至三職等職稱員額', '列委任且職等上限≤3之員額', str(low13), f'≥{need}', '符合' if low13 >= need else '不符', '' if low13 >= need else f'需增 {need - low13} 人',
                         '115.9.30 銓敘部已提修正草案擬刪除本項，發布後本規則停用'))
    else:
        F.append(finding('R4', '第7條第1項', '委任第一至三職等職稱員額', '—', str(low13), '員額未達30人不適用', '不適用'))
    # ---- R5 委四至五得列薦六 ≥ 1/12（第7條第2項）
    w45 = 0; w35 = 0; base = 0; de6 = 0
    for r in rows:
        if not r['拆算']['計入']:
            continue
        g = grades_of(r['職等']); guan = (r['官等'] or '').strip()
        if guan == '委任' and g == {4, 5}:
            w45 += r['員額']
            for m in DE_LIE.finditer(r['備註'] or ''):
                if m.group(2) == '薦任':
                    de6 += c2n(m.group(1))
        if guan == '委任' and g == {3, 4, 5}:
            w35 += r['員額']
        if guan == '委任或薦任' and 5 in g:
            base += r['員額']
    if total >= 12:
        need = math.ceil(base / 12) if base else 0
        ok = w45 >= need
        F.append(finding('R5', '第7條第2項', '委任第四至五職等職稱員額', f'列委四至五之職稱員額 {w45}（內 {de6} 人得列薦六）；委五或薦六至七職稱員額 {base}×1/12', str(w45), f'≥{need}', '符合' if ok else '不符',
                         '' if ok else f'需增 {need - w45} 人', ('另置委三至五職稱 ' + str(w35) + ' 人（組織結構特殊者之替代方式）' if w35 else '')))
    # ---- R6 未滿12人
    if 4 <= total < 12:
        has_w = W > 0
        F.append(finding('R6', '第7條第3項', '未滿12人應置委任職稱', '委任員額', str(W), '≥1', '符合' if has_w else '不符'))
    # ---- 主管判定
    for r in rows:
        t = r['職稱']
        r['主管'] = bool(MGR_RE.search(t)) and not re.search(r'主任管理員|主任管制員|主任報務員|主任航務員|主任氣象員|管理員$|護理長', t)
    # ---- R7 薦七以下薦任非主管（第8條）
    j_non = sum(r['拆算']['薦任'] for r in rows if r['拆算']['計入'] and not r['主管'])
    j_low = sum(r['拆算']['薦任'] for r in rows if r['拆算']['計入'] and not r['主管'] and grades_of(r['職等']) and max(grades_of(r['職等'])) <= 7)
    if total >= 20:
        if level in ('1', '2'):
            need = 40
        else:
            need = 45 if '工程' in category else (5 if '調查' in category else 50)
        p = pct(j_low, j_non)
        ok = p is not None and p >= need
        gap = '' if ok or p is None else f'薦七以下薦任非主管需增 {math.ceil(need / 100 * j_non) - j_low} 人'
        F.append(finding('R7', '第8條', '薦任第七職等以下薦任非主管比率', f'薦七以下薦任非主管 {j_low} ÷ 薦任非主管 {j_non}', f'{p}%' if p is not None else '—', f'≥{need}%', ('符合' if ok else '不符') if p is not None else '需確認', gap,
                         '主管／非主管以職稱判定（…長、主任），請核對'))
    # ---- R3a 非主管不得高於副首長、幕僚長
    staff_heads = [r for r in rows if r['職稱'] in STAFF_HEAD]
    head_title = rows[0]['職稱'] if rows else ''
    deputy = [r for r in rows if r['職稱'] in ('副' + head_title, '政務次長', '常務次長', '副主任委員', '副人事長', '副主計長', '副總裁')]
    cap_rows = staff_heads + deputy
    if cap_rows:
        cap = max((max(grades_of(r['職等'])) for r in cap_rows if grades_of(r['職等'])), default=None)
        over = [r['職稱'] for r in rows if not r['主管'] and r['拆算']['計入'] and grades_of(r['職等']) and cap and max(grades_of(r['職等'])) > cap]
        F.append(finding('R3a', '第4條第1項第4款', '非主管職稱不得高於副首長、幕僚長', f'副首長／幕僚長最高職等 {cap}', '；'.join(sorted(set(over))) or '無', f'≤{cap}', '符合' if not over else '不符', '', '例外需經考試院同意（第13條）'))
    # ---- R3b 接續不跳空
    covered = set(); top = set()
    head_like = set(id(r) for r in rows[:1]) | set(id(r) for r in rows if (r['職稱'].startswith('副') and HEAD_TITLES.match(r['職稱'][1:])) or re.search(r'次長$|副主任委員$|副首長', r['職稱']))
    for r in rows:
        if r['拆算']['計入']:
            (top if id(r) in head_like else covered).__ior__(grades_of(r['職等']))
    if covered:
        lo, hi = min(covered), max(covered)
        gaps = [g for g in range(lo, hi + 1) if g not in covered]
        F.append(finding('R3b', '第4條第1項第6款', '官等職等接續不跳空', f'首長、副首長以外職稱涵蓋第{lo}至第{hi}職等', '缺：' + '、'.join(map(str, gaps)) if gaps else '無跳空', '無跳空', '符合' if not gaps else '不符', '',
                         '首長、副首長（多為組織法律所定）不計；例外需經考試院同意（第13條）'))
    # ---- R9 三級署局 參議
    if level == '3' and category == '中央三級_署局及相當中央三級獨立機關':
        canyi = sum(r['員額'] for r in rows if r['職稱'] == '參議')
        ms = sum(r['員額'] for r in rows if r['職稱'] in STAFF_HEAD)
        first = sum(r['員額'] for r in rows if r['職稱'] in ('組長', '處長') and not (r.get('單位') and re.search(r'人事|主計|會計|政風|秘書', r['單位'] or '')))
        cap9 = math.floor((ms + first) / 2)
        if canyi:
            F.append(finding('R9', '第9條第2款', '參議員額上限', f'（幕僚長 {ms}＋一級業務單位主管 {first}）÷2', str(canyi), f'≤{cap9}（另派出單位主管及所屬機關首長每滿5人＋1，未計）', '符合' if canyi <= cap9 else '需確認', '', '派出單位及所屬機關首長數需人工補充'))
    # ---- R10 簡任非主管 ≤ 二級單位主管×7/10（中央二、三級）
    if level in ('2', '3'):
        k_non = sum(r['拆算']['簡任'] for r in rows if r['拆算']['計入'] and not r['主管'] and r['職稱'] not in CANYI and not re.search(r'研究員|公職律師|消費者保護官', r['職稱']))
        sec = sum(r['員額'] for r in rows if r['職稱'] in SECOND_LEVEL_MGR[level])
        if sec:
            cap10 = math.floor(sec * 0.7)
            ok = k_non <= cap10
            F.append(finding('R10', '第9條第5款', '一、二級單位主管間之簡任非主管員額', f'簡任非主管（不含參事、技監、顧問、參議、研究委員、研究員等）{k_non}；二級單位主管（科長）{sec}×7/10', str(k_non), f'≤{cap10}', '符合' if ok else '不符', '' if ok else f'需減 {k_non - cap10} 人',
                             '修正前原即不符者，修編時比例維持原有或更低者始同意核備（修正總說明）'))
        else:
            F.append(finding('R10', '第9條第5款', '一、二級單位主管間之簡任非主管員額', '—', str(k_non), '—', '需確認', '', '查無科長職稱，二級單位主管數需人工確認'))
    # ---- R11 職稱列等實證對照
    tg = rules['titles'].get(level, {})
    st = collections.Counter(); odd = []
    for r in rows:
        code = grade_code(r['官等'], r['職等'])
        r['列等碼'] = code
        info = tg.get(r['職稱'])
        if r['員額'] <= 0:
            r['列等狀態'] = '兼任/零員額'
        elif not r['拆算']['計入']:
            r['列等狀態'] = '非一般文職（特任、比照、警察、軍職、師級等）'
        elif not info:
            r['列等狀態'] = '本層級未見此職稱'; odd.append(f"{r['職稱']}（{code}）")
        elif code == info['主要列等']:
            r['列等狀態'] = f"主流（{info['列等組合'][0]['機關數']}機關）"
        elif any(c['碼'] == code for c in info['列等組合']):
            n_ = next(c['機關數'] for c in info['列等組合'] if c['碼'] == code)
            r['列等狀態'] = f"少數（{n_}機關；主流 {info['主要列等']}）"
        else:
            r['列等狀態'] = f"實證未見此列等（主流 {info['主要列等']}）"; odd.append(f"{r['職稱']}（{code}；主流 {info['主要列等']}）")
        st[r['列等狀態'].split('（')[0]] += 1
    F.append(finding('R11', '第4條、職務列等表（實證）', '職稱與列等對照同層級實證分布', '逐列比對 rules/title_grades.json', '；'.join(f'{k}{v}' for k, v in st.most_common()), '—', '符合' if not odd else '需確認', '', '；'.join(odd[:15]) + ('…' if len(odd) > 15 else '')))
    # ---- 彙總
    bad = [f for f in F if f['結果'].startswith('不符')]
    chk = [f for f in F if f['結果'] == '需確認']
    return {'機關': a['機關'], '層級': level, '隸屬': a['隸屬'], '部會': a['部會'], '類別': category, '類別已確認': confirmed, '雙軌制': dual,
            '編制表': a.get('編制表來源'), '驗算狀態': a.get('驗算狀態'), '員額合計': a.get('員額合計'), '計算': counts,
            '簡任比率': sj, '委任比率': sw, '簡任上限': (th or {}).get('簡任') if th else None, '委任下限': (th or {}).get('委任') if th else None,
            '不符數': len(bad), '需確認數': len(chk), '不符規則': '、'.join(f['規則'] for f in bad), '規則': F,
            '明細': [{k: r.get(k) for k in ('序', '單位', '職稱', '官等', '職等', '員額', '備註', '拆算', '主管', '列等碼', '列等狀態')} for r in rows]}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--master', default=os.path.join(ROOT, 'org-compare', 'data', '組織法編制表_總表_20261031.xlsx'))
    ap.add_argument('--out', default=os.path.join(HERE, 'output'))
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    rules = load_rules()
    import openpyxl
    wb = openpyxl.load_workbook(a.master, read_only=True)
    ov = list(wb['機關總覽'].iter_rows(values_only=True)); oi = {h: i for i, h in enumerate(ov[0])}
    agencies = [{h: r[i] for h, i in oi.items()} for r in ov[1:]]
    for ag in agencies:
        ag['層級'] = str(ag['層級'])
    det = list(wb['編制表明細'].iter_rows(values_only=True)); di = {h: i for i, h in enumerate(det[0])}
    by = collections.defaultdict(list)
    for r in det[1:]:
        try:
            n = int(r[di['員額']])
        except (TypeError, ValueError):
            n = 0
        by[r[di['機關名稱']]].append({'序': r[di['序']], '單位': r[di['單位']], '職稱': (r[di['職稱']] or '').strip(), '官等': r[di['官等']], '職等': r[di['職等']], '員額': n, '備註': r[di['備註']] or ''})
    cats = load_or_make_categories(agencies)
    results = []
    for ag in agencies:
        rows = by.get(ag['機關'], [])
        if not rows:
            continue
        results.append(check_agency(ag, rows, cats[ag['機關']], rules))
    meta = {'產生': dt.datetime.now().strftime('%Y-%m-%d %H:%M'), '總表': os.path.basename(a.master), '機關數': len(results),
            '準則版本': '114.12.4 修正', '規則清單': [l.strip() for l in __doc__.splitlines() if re.match(r'\s+R\d', l)]}
    json.dump({'meta': meta, 'results': results}, open(os.path.join(a.out, 'check_results.json'), 'w', encoding='utf-8'), ensure_ascii=False)
    with open(os.path.join(a.out, '檢核結果_全機關.csv'), 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f)
        w.writerow(['機關', '層級', '隸屬', '類別', '類別已確認', '定有官等員額', '委任', '薦任', '簡任', '簡任比率%', '簡任上限%', '委任比率%', '委任下限%', '不符數', '不符規則', '需確認數'] + ['R1', 'R2', 'R4', 'R5', 'R7', 'R10', 'R3a', 'R3b', 'R11', 'R13'])
        for r in results:
            res = {f['規則']: f['結果'] for f in r['規則']}
            w.writerow([r['機關'], r['層級'], r['隸屬'], r['類別'], '是' if r['類別已確認'] else '', r['計算']['定有官等職等員額'], r['計算']['委任'], r['計算']['薦任'], r['計算']['簡任'],
                        r['簡任比率'], r['簡任上限'], r['委任比率'], r['委任下限'], r['不符數'], r['不符規則'], r['需確認數']] + [res.get(k, '') for k in ['R1', 'R2', 'R4', 'R5', 'R7', 'R10', 'R3a', 'R3b', 'R11', 'R13']])
    c = collections.Counter()
    for r in results:
        for f in r['規則']:
            c[(f['規則'], f['結果'].split(' /')[0])] += 1
    print(f'檢核 {len(results)} 機關 → {a.out}/check_results.json、檢核結果_全機關.csv')
    for rule in ['R1', 'R2', 'R3a', 'R3b', 'R4', 'R5', 'R6', 'R7', 'R9', 'R10', 'R11', 'R13']:
        row = {k[1]: v for k, v in c.items() if k[0] == rule}
        if row:
            print(f'  {rule}: ' + '、'.join(f'{k}{v}' for k, v in sorted(row.items(), key=lambda kv: -kv[1])))
    return 0


if __name__ == '__main__':
    sys.exit(main())
